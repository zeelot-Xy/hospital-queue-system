from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE_CHAPTER = ROOT / "joe documentation" / "CHAPTER ONE with Correction.docx"
CLEAN_CHAPTER = ROOT / "joe documentation" / "CHAPTER ONE - CLEAN.docx"
REPORT_MARKDOWN = ROOT / "docs" / "PROJECT_REPORT.md"
FINAL_REPORT = ROOT / "docs" / "FINAL_PROJECT_REPORT.docx"


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def clean_chapter_one():
    doc = Document(SOURCE_CHAPTER)

    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text == "1.7 Definition of Key Terms":
            paragraph.text = "1.6 Definition of Key Terms"
            text = paragraph.text

        if text == "Note" or text.startswith("Weldone!") or text.startswith("Kindly check") or text.startswith("Let me know") or text.startswith("Effect the corrections"):
            remove_paragraph(paragraph)
            continue

        duplicate = "Many clinics continue to rely on manual methods for managing patient appointments and queues. These traditional approaches present several operational challenges, including:"
        if text.startswith(duplicate) and text.count(duplicate) > 1:
            paragraph.text = text[text.find(duplicate, len(duplicate)) :]
            text = paragraph.text.strip()

        is_heading = (
            text in {"CHAPTER ONE", "INTRODUCTION"}
            or bool(re.match(r"^1\.\d+\s", text))
        )

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if text in {"CHAPTER ONE", "INTRODUCTION"} else WD_ALIGN_PARAGRAPH.JUSTIFY
        )
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.first_line_indent = None if is_heading else Inches(0.5)

        for run in paragraph.runs:
            set_font(run, size=12, bold=is_heading)

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    doc.save(CLEAN_CHAPTER)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_begin, instr_text, fld_char_end])
    set_font(run, size=10)


def configure_report_styles(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for style_name, size, before, after in [
        ("Heading 1", 14, 12, 8),
        ("Heading 2", 13, 10, 6),
        ("Heading 3", 12, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)

    add_page_number(section.footer.paragraphs[0])

    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def add_title_page(doc):
    for _ in range(5):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(24)
    run = title.add_run("DESIGN AND IMPLEMENTATION OF A SMART CLINIC\nAPPOINTMENT AND PATIENT QUEUE MANAGEMENT SYSTEM")
    set_font(run, size=16, bold=True)

    byline = doc.add_paragraph()
    byline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    byline.paragraph_format.space_after = Pt(18)
    set_font(byline.add_run("BY"), size=12, bold=True)

    identity = doc.add_paragraph()
    identity.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(identity.add_run("[STUDENT NAME]\n[MATRICULATION NUMBER]"), size=12, bold=True)

    for _ in range(3):
        doc.add_paragraph()

    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.line_spacing = 1.5
    set_font(
        statement.add_run(
            "A FINAL-YEAR PROJECT SUBMITTED TO THE DEPARTMENT OF [ACADEMIC DEPARTMENT], "
            "[INSTITUTION NAME], IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF [DEGREE]"
        ),
        size=12,
    )

    for _ in range(3):
        doc.add_paragraph()
    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(date.add_run("[MONTH YEAR]"), size=12, bold=True)
    doc.add_page_break()


def add_toc(doc):
    heading = doc.add_paragraph(style="Heading 1")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("TABLE OF CONTENTS")
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and update this field in Microsoft Word."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, placeholder, end])
    set_font(run, size=12)
    doc.add_page_break()


def parse_table(lines, start):
    rows = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def add_table(doc, rows):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    usable_width_dxa = 8670
    base_width = usable_width_dxa // columns
    column_widths = [base_width] * columns
    column_widths[-1] += usable_width_dxa - sum(column_widths)

    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(usable_width_dxa))
    table_indent = OxmlElement("w:tblInd")
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), "120")
    table_properties.append(table_indent)

    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, column_widths):
        grid_column.set(qn("w:w"), str(width))

    for row_index, values in enumerate(rows):
        row = table.add_row()
        cells = row.cells
        if row_index == 0:
            table_row_properties = row._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            table_row_properties.append(header)
        for column_index, cell in enumerate(cells):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(column_widths[column_index]))
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            text = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(text)
            set_font(run, size=10, bold=row_index == 0)
            if row_index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7E6E6")
                cell._tc.get_or_add_tcPr().append(shade)
    doc.add_paragraph()


def add_markdown_content(doc):
    lines = REPORT_MARKDOWN.read_text(encoding="utf-8").splitlines()
    index = 0
    skipping_title_metadata = True

    while index < len(lines):
        raw = lines[index].rstrip()
        text = raw.strip()

        if text == "## Abstract":
            skipping_title_metadata = False

        if skipping_title_metadata:
            index += 1
            continue

        if not text or text.startswith(">"):
            index += 1
            continue

        if text.startswith("|"):
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        if text.startswith("# "):
            if text not in {"# Chapter One: Introduction"}:
                doc.add_page_break()
            paragraph = doc.add_paragraph(style="Heading 1")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.add_run(text[2:].upper())
        elif text.startswith("## "):
            paragraph = doc.add_paragraph(text[3:], style="Heading 2")
        elif text.startswith("### "):
            paragraph = doc.add_paragraph(text[4:], style="Heading 3")
        elif re.match(r"^\d+\.\s", text):
            content = re.sub(r"^\d+\.\s", "", text)
            paragraph = doc.add_paragraph(style="List Number")
            add_inline_markdown(paragraph, content)
        elif text.startswith("- [ ] "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, "[ ] " + text[6:])
        elif text.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_markdown(paragraph, text[2:])
        else:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.first_line_indent = Inches(0.5)
            add_inline_markdown(paragraph, text)
        index += 1


def add_inline_markdown(paragraph, text):
    parts = re.split(r"(\*\*.*?\*\*|\*.*?\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=12, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=12, italic=True)
        else:
            run = paragraph.add_run(part.replace("`", ""))
            set_font(run, size=12)


def build_final_report():
    doc = Document()
    configure_report_styles(doc)
    add_title_page(doc)
    add_toc(doc)
    add_markdown_content(doc)
    doc.core_properties.title = "Design and Implementation of a Smart Clinic Appointment and Patient Queue Management System"
    doc.core_properties.subject = "Final-Year Project Report"
    doc.core_properties.author = "[Student Name]"
    doc.save(FINAL_REPORT)


if __name__ == "__main__":
    clean_chapter_one()
    build_final_report()
    print(f"Created {CLEAN_CHAPTER}")
    print(f"Created {FINAL_REPORT}")
