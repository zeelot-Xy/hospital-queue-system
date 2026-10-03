from __future__ import annotations

from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "joe documentation" / "CHAPTER TWO - LITERATURE REVIEW.docx"
MARKDOWN = ROOT / "docs" / "CHAPTER_THREE_METHODOLOGY_AND_SYSTEM_DESIGN.md"
OUTPUT = ROOT / "joe documentation" / "CHAPTER THREE - RESEARCH METHODOLOGY AND SYSTEM DESIGN.docx"


def font_run(run, *, bold=None, italic=None, size=12):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rpr.rFonts.set(qn(key), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin, section.right_margin = Inches(1.25), Inches(1)
    section.top_margin, section.bottom_margin = Inches(1), Inches(1)
    for name in ("Normal", "Normal (Web)"):
        if name not in doc.styles:
            continue
        st = doc.styles[name]
        st.font.name, st.font.size = "Times New Roman", Pt(12)
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        pf = st.paragraph_format
        pf.alignment, pf.line_spacing, pf.space_after = WD_ALIGN_PARAGRAPH.JUSTIFY, 1.5, Pt(0)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold = "Times New Roman", Pt(12), True
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        pf = st.paragraph_format
        pf.line_spacing, pf.space_before, pf.space_after, pf.keep_with_next = 1.5, Pt(0), Pt(0), True


def add_page_number_field(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in list(p.runs):
        p._p.remove(run._r)
    run = p.add_run()
    font_run(run, size=10)
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_inline(p, text):
    for part in re.split(r"(\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("*") and part.endswith("*"):
            font_run(p.add_run(part[1:-1]), italic=True)
        else:
            font_run(p.add_run(part))


def body(doc, text):
    p = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing, pf.space_before, pf.space_after = 1.5, Pt(0), Pt(0)
    pf.first_line_indent = Inches(0.5)
    add_inline(p, text)


def heading(doc, text, level, centered=False):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    font_run(p.add_run(text), bold=True)


def numbered_item(doc, text):
    p = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.first_line_indent = Inches(-0.3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text)


def reference(doc, text):
    p = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    add_inline(p, text)


def set_cell_shading(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:w"), str(dxa)); tcw.set(qn("w:type"), "dxa")


def table(doc, rows):
    count = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=count)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.style = "Table Grid"
    widths = {2:[2700,5970],3:[1300,3650,3720],4:[1800,2500,2600,1770]}.get(count,[8670//count]*count)
    for r_index, values in enumerate(rows):
        for c_index, value in enumerate(values):
            cell = t.cell(r_index,c_index)
            set_cell_width(cell,widths[c_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if r_index == 0:
                set_cell_shading(cell,"D9EAF4")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            font_run(p.add_run(value),bold=(r_index==0),size=10)
    # repeat header row
    trpr = t.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trpr.append(tbl_header)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def figure(doc, path, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(5.95))
    docpr = shape._inline.docPr
    docpr.set("name", alt[:80])
    docpr.set("descr", alt)


def caption(doc, text):
    p = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    font_run(p.add_run(text), italic=True, size=11)


def table_title(doc, text):
    p = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    font_run(p.add_run(text), bold=True, size=11)


def build():
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc); configure(doc); add_page_number_field(doc)
    lines = MARKDOWN.read_text(encoding="utf-8").splitlines()
    in_refs = False
    i = 0
    while i < len(lines):
        text = lines[i].strip()
        if not text:
            i += 1; continue
        if text == "# CHAPTER THREE":
            heading(doc,"CHAPTER THREE",1,True)
        elif text == "## RESEARCH METHODOLOGY AND SYSTEM DESIGN":
            heading(doc,"RESEARCH METHODOLOGY AND SYSTEM DESIGN",2,True)
        elif text == "## References":
            doc.add_page_break(); heading(doc,"REFERENCES",1,True); in_refs=True
        elif text.startswith("#### "):
            heading(doc,text[5:],3)
        elif text.startswith("### "):
            heading(doc,text[4:],2)
        elif text.startswith("!["):
            match=re.match(r"!\[(.+?)\]\((.+?)\)",text)
            if match:
                figure(doc,ROOT/"docs"/match.group(2),match.group(1))
        elif text.startswith("Figure 3."):
            caption(doc,text)
        elif text.startswith("Table 3."):
            table_title(doc,text)
        elif text.startswith("|"):
            raw=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                raw.append([c.strip() for c in lines[i].strip().strip("|").split("|")]); i+=1
            data=[row for row in raw if not all(re.fullmatch(r":?-{3,}:?",c) for c in row)]
            table(doc,data); continue
        elif re.match(r"^\d+\. ",text):
            numbered_item(doc,text)
        elif in_refs:
            reference(doc,text)
        else:
            body(doc,text)
        i += 1
    doc.core_properties.title = "Chapter Three - Research Methodology and System Design"
    doc.core_properties.subject = "Smart Clinic Appointment and Patient Queue Management System"
    doc.core_properties.author = "Smart Clinic Project"
    doc.core_properties.comments = "Chapter Three prepared from the implemented system and timestamped verification evidence."
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
