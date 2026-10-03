from __future__ import annotations

from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "joe documentation" / "CHAPTER ONE - CLEAN.docx"
MARKDOWN = ROOT / "docs" / "CHAPTER_TWO_LITERATURE_REVIEW.md"
OUTPUT = ROOT / "joe documentation" / "CHAPTER TWO - LITERATURE REVIEW.docx"


def set_run_font(run, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Times New Roman")
    rpr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rpr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure_styles(doc: Document) -> None:
    for name in ("Normal", "Normal (Web)"):
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(0)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True


def add_inline(paragraph, text: str) -> None:
    parts = re.split(r"(\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("*") and part.endswith("*"):
            set_run_font(paragraph.add_run(part[1:-1]), italic=True)
        else:
            set_run_font(paragraph.add_run(part))


def add_body(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    add_inline(paragraph, text)


def add_heading(doc: Document, text: str, level: int, *, centered: bool = False) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = None
    run = paragraph.add_run(text)
    set_run_font(run, bold=True)


def add_reference(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Normal (Web)" if "Normal (Web)" in doc.styles else "Normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    add_inline(paragraph, text)


def build() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_styles(doc)

    in_references = False
    for raw in MARKDOWN.read_text(encoding="utf-8").splitlines():
        text = raw.strip()
        if not text:
            continue
        if text == "# CHAPTER TWO":
            add_heading(doc, "CHAPTER TWO", 1, centered=True)
        elif text == "## LITERATURE REVIEW":
            add_heading(doc, "LITERATURE REVIEW", 2, centered=True)
        elif text == "## References":
            doc.add_page_break()
            add_heading(doc, "REFERENCES", 1, centered=True)
            in_references = True
        elif text.startswith("#### "):
            add_heading(doc, text[5:], 3)
        elif text.startswith("### "):
            add_heading(doc, text[4:], 2)
        elif in_references:
            add_reference(doc, text)
        else:
            add_body(doc, text)

    doc.core_properties.title = "Chapter Two - Literature Review"
    doc.core_properties.subject = "Smart Clinic Appointment and Patient Queue Management System"
    doc.core_properties.comments = "Verified Chapter Two revision with current literature through 2026."
    doc.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build()
