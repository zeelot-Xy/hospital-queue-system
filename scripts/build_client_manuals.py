from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "client-manuals"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = RGBColor(11, 37, 69)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 100, 115)
LIGHT = "E8EEF5"
WHITE = RGBColor(255, 255, 255)


def set_font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    first_row_props = table.rows[0]._tr.get_or_add_trPr()
    header_marker = first_row_props.find(qn("w:tblHeader"))
    if header_marker is None:
        header_marker = OxmlElement("w:tblHeader")
        header_marker.set(qn("w:val"), "true")
        first_row_props.append(header_marker)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    values = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, RGBColor(31, 77, 120), 10, 5),
    }
    for name, (size, color, before, after) in values.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_header_footer(doc, short_title):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run(short_title)
    set_font(run, size=9, color=MUTED, bold=True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Smart Clinic Appointment and Patient Queue Management System | Client Manual | Version 1.0")
    set_font(run, size=8.5, color=MUTED)


def add_cover(doc, title, subtitle, audience):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(70)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("CLIENT OPERATIONS GUIDE")
    set_font(run, size=11, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(title)
    set_font(run, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    run = p.add_run(subtitle)
    set_font(run, size=14, color=MUTED)

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_geometry(table, [2600, 6760])
    labels = [("WHO THIS IS FOR", audience), ("SYSTEM VERSION", "1.0"), ("SUPPORT NOTE", "Keep this guide beside the clinic host computer or server records.")]
    for row_index, (label, value) in enumerate(labels):
        if row_index:
            cells = table.add_row().cells
            set_table_geometry(table, [2600, 6760])
        else:
            cells = table.rows[0].cells
        shade(cells[0], LIGHT)
        r = cells[0].paragraphs[0].add_run(label)
        set_font(r, size=9, color=NAVY, bold=True)
        r = cells[1].paragraphs[0].add_run(value)
        set_font(r, size=10.5)
    doc.add_page_break()


def add_picture(doc, path, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("descr", caption)
    picture._inline.docPr.set("title", caption)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(caption), size=9, color=MUTED, italic=True)


def add_steps(doc, title, steps):
    doc.add_heading(title, level=2)
    numbering = doc.part.numbering_part.element
    existing_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    num_id = max(existing_ids or [0]) + 1
    style_num_id = doc.styles["List Number"]._element.pPr.numPr.numId.val
    style_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == style_num_id
    )
    abstract_id = style_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for heading, detail in steps:
        p = doc.add_paragraph(style="List Number")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = 0
        num_pr.get_or_add_numId().val = num_id
        r = p.add_run(heading + ". ")
        set_font(r, bold=True)
        r = p.add_run(detail)
        set_font(r)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(item))


def add_quick_flow(doc, labels):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    set_table_geometry(table, [2340, 2340, 2340, 2340])
    for index, label in enumerate(labels):
        cell = table.rows[0].cells[index]
        shade(cell, LIGHT)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"STEP {index + 1}\n{label}")
        set_font(r, size=10, color=NAVY, bold=True)


def local_manual():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc, "Clinic Computer and Local Network Manual")
    add_cover(
        doc,
        "Clinic Computer and Local Network Manual",
        "Simple installation, daily operation, backup, and recovery instructions",
        "Clinic administrators and the person responsible for the main clinic computer",
    )

    doc.add_heading("1. What this package does", level=1)
    doc.add_paragraph("One Windows computer becomes the clinic host. Staff, doctors, and authorized users open the system in a web browser using the host computer's local network address. Patient and clinic data stays on the clinic host unless the clinic creates an external backup.")
    add_quick_flow(doc, ["Prepare", "Install", "Open", "Back up"])

    doc.add_heading("2. Before you begin", level=1)
    add_bullets(doc, [
        "A Windows 10 or Windows 11 computer that can remain switched on during clinic hours.",
        "At least 8 GB RAM, 20 GB free disk space, and a stable clinic router or Wi-Fi network.",
        "Docker Desktop installed and running on the host computer.",
        "The clinic computer's LAN address, such as 192.168.1.20.",
        "A strong administrator password stored privately by the clinic owner or manager.",
    ])

    add_steps(doc, "3. First installation", [
        ("Extract the package", "Place the extracted folder in a permanent location such as C:\\ClinicSystem. Do not run it from a USB drive"),
        ("Start Docker Desktop", "Wait until Docker Desktop reports that the engine is running"),
        ("Run guided installation", "Open deploy\\local and double-click INSTALL-CLINIC.cmd. Enter the administrator name, email, phone, and a private password when requested"),
        ("Keep settings private", "Database and security secrets are generated automatically. Do not share the deploy\\local\\.env file"),
        ("Open the application", "On the host computer open http://localhost:8080. On another clinic device open the LAN address followed by :8080"),
        ("Sign in", "Use the administrator email and password entered in the .env file, then immediately store them in the clinic's secure password record"),
    ])

    doc.add_heading("4. Normal daily operation", level=1)
    doc.add_heading("Start and stop", level=2)
    add_bullets(doc, [
        "Start Docker Desktop, then double-click START-CLINIC.cmd or OPEN-CLINIC.cmd.",
        "Check the system with: .\\clinic.ps1 status.",
        "At the end of the day, double-click STOP-CLINIC.cmd before shutting down the host computer.",
    ])
    doc.add_heading("Role workflow", level=2)
    add_bullets(doc, [
        "Administrator/staff: manage departments, assign doctor accounts, register walk-ins, monitor queues, confirm admission, and view reports.",
        "Doctor: set availability, call the next patient, start consultation, record findings, and complete the visit.",
        "Patient: register, book an available appointment, mark arrival near the appointment time, follow queue updates, and view visit history.",
    ])

    add_steps(doc, "5. Create a backup", [
        ("Keep the clinic system running", "Docker Desktop and the database must be available"),
        ("Run the backup command", "In PowerShell run: .\\clinic.ps1 backup"),
        ("Copy the backup elsewhere", "Copy the new SQL file from the backups folder to an encrypted USB drive or other clinic-approved secure location"),
        ("Record the date", "Write the backup date in the clinic's backup register"),
    ])

    add_steps(doc, "6. Restore after a problem", [
        ("Protect the current files", "Do not delete the current system folder. Make an extra copy before restoring"),
        ("Choose the correct backup", "Use the newest known-good SQL backup"),
        ("Run restore", "Run: .\\clinic.ps1 restore -BackupFile C:\\Path\\To\\hospital-queue-backup.sql"),
        ("Verify", "Sign in and confirm that departments, users, appointments, queues, and reports are present"),
    ])

    doc.add_heading("7. Update the system", level=1)
    add_bullets(doc, [
        "Create and copy a backup before every update.",
        "Replace the application files only with the approved new release package. Keep the existing deploy\\local\\.env file.",
        "Run .\\clinic.ps1 update and wait for the services to become healthy.",
        "Sign in and test one administrator page, one doctor page, and one patient page.",
    ])

    doc.add_heading("8. Troubleshooting", level=1)
    add_bullets(doc, [
        "Page does not open: confirm Docker Desktop is running and run .\\clinic.ps1 status.",
        "Other devices cannot connect: confirm they are on the same network, use the correct LAN address, and allow port 8080 through Windows Firewall.",
        "Login fails: check the exact email, password, and Caps Lock. Do not edit database files manually.",
        "Service is unhealthy: run .\\clinic.ps1 stop, restart Docker Desktop manually, then run .\\clinic.ps1 start.",
        "Data appears missing: stop work and contact the technical maintainer before restoring or deleting anything.",
    ])

    doc.add_heading("9. Security and emergency rules", level=1)
    add_bullets(doc, [
        "Never send the .env file, database backup, or administrator password through ordinary chat or email.",
        "Do not expose port 8080 directly to the public internet.",
        "Keep Windows, Docker Desktop, and browsers updated.",
        "If the host computer is stolen or infected, disconnect it from the network and contact the clinic's technical maintainer.",
        "Keep at least one recent encrypted backup away from the clinic computer.",
    ])
    doc.save(OUT / "Clinic_Computer_and_Local_Network_Manual.docx")


def online_manual():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc, "Online Server Deployment Manual")
    add_cover(
        doc,
        "Online Server Deployment Manual",
        "Generic Ubuntu VPS installation, HTTPS, backup, update, and recovery",
        "The client's server administrator or contracted technical maintainer",
    )

    doc.add_heading("1. What this package does", level=1)
    doc.add_paragraph("This package runs the complete clinic system on a client-chosen Ubuntu server. Caddy provides HTTPS automatically, the frontend and API run in separate containers, and PostgreSQL data remains in a persistent server volume.")
    add_quick_flow(doc, ["Prepare VPS", "Point domain", "Deploy", "Verify"])

    doc.add_heading("2. Server requirements", level=1)
    add_bullets(doc, [
        "Ubuntu 22.04 or 24.04 LTS with at least 2 CPU cores, 4 GB RAM, and 30 GB storage.",
        "A domain name whose DNS A record points to the server's public IP address.",
        "Inbound ports 22, 80, and 443 allowed by the provider firewall.",
        "Docker Engine and the Docker Compose plugin installed.",
        "SSH access using a non-root sudo account and preferably an SSH key.",
    ])

    add_steps(doc, "3. First deployment", [
        ("Prepare DNS", "Create the domain A record and wait until it resolves to the server"),
        ("Upload and extract", "Upload the approved online ZIP package, extract it to /opt/hospital-queue, and restrict access to the server administrator"),
        ("Create settings", "In the deploy folder copy .env.example to .env and replace every CHANGE_TO value"),
        ("Set the domain", "Enter the real DOMAIN and ACME_EMAIL values. Do not include http:// or https:// in DOMAIN"),
        ("Protect the settings", "Run chmod 600 .env and chmod +x clinic.sh"),
        ("Install", "Run sudo ./clinic.sh install and wait for all services to become healthy"),
        ("Verify HTTPS", "Open https://your-domain and confirm the browser shows a valid secure connection"),
        ("Sign in", "Use the administrator credentials stored in .env and complete the initial doctor and department setup"),
    ])

    doc.add_heading("4. Routine commands", level=1)
    add_bullets(doc, [
        "View status: sudo ./clinic.sh status",
        "Start: sudo ./clinic.sh start",
        "Stop: sudo ./clinic.sh stop",
        "Create backup: sudo ./clinic.sh backup",
        "Install an approved update: sudo ./clinic.sh update",
    ])

    add_steps(doc, "5. Backup and off-server storage", [
        ("Create a database dump", "Run sudo ./clinic.sh backup"),
        ("Confirm the file", "Check that the new SQL file exists and is not empty"),
        ("Encrypt and copy", "Copy the backup to client-approved encrypted storage outside the VPS"),
        ("Retain multiple dates", "Keep daily recent backups and at least one older monthly recovery point according to clinic policy"),
        ("Test restoration", "Use a controlled maintenance window to rehearse restore on a separate test deployment"),
    ])

    add_steps(doc, "6. Restore procedure", [
        ("Declare maintenance", "Prevent staff from changing data during restoration"),
        ("Take a safety backup", "Back up the current database even if it may contain incorrect data"),
        ("Restore", "Run sudo ./clinic.sh restore /secure/path/to/backup.sql"),
        ("Restart and verify", "Start the services and verify administrator login, recent appointments, queue history, and reports"),
        ("Record the incident", "Document the backup used, operator, date, result, and any lost time period"),
    ])

    doc.add_heading("7. Safe updates", level=1)
    add_bullets(doc, [
        "Read the release notes and create an off-server backup first.",
        "Keep the current .env file; never replace it with .env.example.",
        "Upload the new approved application files and run sudo ./clinic.sh update.",
        "Confirm HTTPS, login, booking, queue status, consultation completion, reports, and Socket.IO updates.",
        "Keep the previous release archive until the update has been accepted.",
    ])

    doc.add_heading("8. Monitoring and troubleshooting", level=1)
    add_bullets(doc, [
        "Domain does not open: check DNS, provider firewall rules, and sudo ./clinic.sh status.",
        "HTTPS certificate is pending: confirm ports 80 and 443 are public and the domain points to this server.",
        "Database is unhealthy: do not remove volumes. Check free disk space and container logs before recovery.",
        "Web page opens but actions fail: check backend health and confirm ALLOWED_ORIGINS is generated from the correct DOMAIN.",
        "Repeated restarts: preserve logs and contact the maintainer; do not repeatedly recreate containers or volumes.",
    ])

    doc.add_heading("9. Security checklist", level=1)
    add_bullets(doc, [
        "Use SSH keys, disable password-based root login, and apply Ubuntu security updates.",
        "Never publish .env, SQL backups, or administrator passwords in source control or support messages.",
        "Allow PostgreSQL only inside the Docker network; do not open port 5432 publicly.",
        "Limit server access to named maintainers and remove access immediately when responsibilities change.",
        "Review backups, disk space, certificate renewal, and dependency advisories regularly.",
    ])

    doc.add_heading("10. Emergency recovery", level=1)
    doc.add_paragraph("If the server is lost, create a replacement Ubuntu VPS, point the domain to it, install Docker, deploy the same approved package with secure settings, restore the newest verified off-server backup, and complete the verification checklist before reopening the system to clinic users.")
    doc.save(OUT / "Online_Server_Deployment_Manual.docx")


def evaluation_manual():
    doc = Document()
    configure_styles(doc)
    add_header_footer(doc, "Client Evaluation and Guided Tour Manual")
    add_cover(doc, "Client Evaluation and Guided Tour Manual", "One-click setup and a safe 15-minute tour through every clinic role", "Clinic owners, project evaluators, and first-time users")

    doc.add_heading("1. What this evaluation edition is", level=1)
    doc.add_paragraph("This separate Windows package lets the client explore the complete clinic workflow without entering real patient data. It uses its own Docker project and database volume, so its demonstration records do not mix with either production delivery.")
    add_quick_flow(doc, ["Install", "Read credentials", "Explore roles", "Reset demo"])
    add_bullets(doc, [
        "Use demonstration information only. The amber banner remains visible while evaluation mode is active.",
        "The installer creates a different random password on each computer and writes it to DEMO-CREDENTIALS.txt.",
        "The reset command replaces only the named demonstration accounts and their related workflow records.",
    ])

    add_steps(doc, "2. Install in three actions", [
        ("Extract the ZIP", "Extract the complete evaluation package to a normal folder; do not run it inside the ZIP"),
        ("Start Docker Desktop", "Wait until Docker Desktop reports that its engine is running"),
        ("Double-click the installer", "Open deploy\\evaluation and run INSTALL-EVALUATION.cmd. The first build may take several minutes"),
    ])
    doc.add_paragraph("When installation finishes, the browser opens http://localhost:8081. Open DEMO-CREDENTIALS.txt in the deploy folder to see the four account emails and generated shared demonstration password.")

    doc.add_heading("3. Recognize evaluation mode", level=1)
    add_picture(doc, ROOT / "evidence" / "release-verification" / "2026-08-11_062512" / "01-login-desktop.png", "Sign-in page. The evaluation build also shows an amber safety banner and Guided Tour button.")
    doc.add_paragraph("Select Guided Tour at any time for role-specific next steps.")

    add_steps(doc, "4. Administrator - about 3 minutes", [
        ("Sign in", "Use admin.demo@clinic.local and the password in DEMO-CREDENTIALS.txt"),
        ("Review Accounts", "Create reception staff, control access, or issue replacement passwords"),
        ("Review setup", "Open Departments and Doctors to see starter departments and the assigned Demo Doctor"),
    ])
    add_picture(doc, ROOT / "evidence" / "release-verification" / "2026-08-11_062512" / "03-staff-dashboard.png", "Staff operations: live queue, appointments, reports, doctors, departments, and administrator account controls.")

    add_steps(doc, "5. Doctor - about 4 minutes", [
        ("Sign in", "Use doctor.demo@clinic.local"),
        ("Review availability", "Availability is prepared for every day so booking can be explored"),
        ("Call the patient", "Select Call Next Patient; the staff view receives a real-time alert"),
    ])
    add_picture(doc, ROOT / "evidence" / "release-verification" / "2026-08-11_062512" / "04-doctor-dashboard.png", "Doctor dashboard: availability, current consultation, notifications, and queue controls.")

    add_steps(doc, "6. Staff - about 4 minutes", [
        ("Sign in", "Use staff.demo@clinic.local"),
        ("Open Live Queue", "Confirm that the called Demo Patient is sent into the consultation room"),
        ("Explore operations", "Review Appointments, Reports, Doctors, Departments, and Add Walk-in"),
    ])
    add_steps(doc, "7. Patient - about 4 minutes", [
        ("Sign in", "Use patient.demo@clinic.local"),
        ("Review Queue Status", "See the demonstration queue number and real-time state"),
        ("Review Visit History", "See the completed sample consultation"),
        ("Explore booking", "Inspect departments, doctors, dates, and available time slots"),
    ])
    add_picture(doc, ROOT / "evidence" / "release-verification" / "2026-08-11_062512" / "05-patient-dashboard-desktop.png", "Patient dashboard: booking, appointments, profile, visit history, notifications, and queue status.", width=5.4)

    add_steps(doc, "8. Restore the original demonstration", [
        ("Sign out", "Close active demonstration sessions"),
        ("Run RESET-DEMO.cmd", "Double-click it in deploy\\evaluation and wait for success"),
        ("Open the clinic", "Run OPEN-CLINIC.cmd and reuse DEMO-CREDENTIALS.txt"),
    ])
    doc.add_paragraph("The reset command is guarded by DEMO_MODE=true and affects only the fixed demonstration account emails.")

    doc.add_heading("9. Everyday controls", level=1)
    add_bullets(doc, [
        "OPEN-CLINIC.cmd - open the evaluation system.",
        "CHECK-CLINIC.cmd - display service status.",
        "RESET-DEMO.cmd - restore the original sample workflow.",
        "STOP-CLINIC.cmd - stop evaluation services.",
    ])
    doc.add_page_break()
    doc.add_heading("10. Troubleshooting and production handoff", level=1)
    add_bullets(doc, [
        "Docker error: start Docker Desktop manually and wait until it is ready.",
        "Page does not open: run CHECK-CLINIC.cmd, then OPEN-CLINIC.cmd.",
        "Workflow is already completed: run RESET-DEMO.cmd.",
        "Do not convert the evaluation database into the live clinic database. Install the chosen production package with new private passwords and an empty production database.",
    ])
    doc.save(OUT / "Client_Evaluation_and_Guided_Tour_Manual.docx")


if __name__ == "__main__":
    local_manual()
    online_manual()
    evaluation_manual()
    print(f"Created manuals in {OUT}")
